"""Lectura documental sin credenciales, red ni archivos institucionales reales."""

from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.drive import cliente, extractores
from app.drive.cliente import ErrorDrive
from app.drive.extractores import ErrorExtraccion, extraer_texto_desde_bytes


class SolicitudFalsa:
    def __init__(self, resultado):
        self.resultado = resultado

    def execute(self):
        if isinstance(self.resultado, Exception):
            raise self.resultado
        return self.resultado


class ServicioFalso:
    def __init__(self, *, paginas=None, metadatos=None, contenido=b"PDF exportado"):
        self.paginas = paginas if paginas is not None else [{"files": []}]
        self.metadatos = metadatos or {
            "id": "documento-ficticio",
            "name": "Institucional.pdf",
            "mimeType": "application/pdf",
            "webViewLink": "https://drive.google.com/file/d/documento-ficticio/view",
        }
        self.contenido = contenido
        self.llamadas = []
        self.cierres = 0

    def files(self):
        return self

    def list(self, **parametros):
        self.llamadas.append(("list", parametros))
        return SolicitudFalsa(self.paginas.pop(0))

    def get(self, **parametros):
        self.llamadas.append(("get", parametros))
        return SolicitudFalsa(self.metadatos)

    def get_media(self, **parametros):
        self.llamadas.append(("get_media", parametros))
        return SolicitudFalsa(self.contenido)

    def export_media(self, **parametros):
        self.llamadas.append(("export_media", parametros))
        return SolicitudFalsa(self.contenido)

    def close(self):
        self.cierres += 1


class DescargaFalsa:
    def __init__(self, destino, solicitud):
        self.destino = destino
        self.contenido = solicitud.resultado
        self.pasos = 0

    def next_chunk(self):
        if isinstance(self.contenido, Exception):
            raise self.contenido
        self.pasos += 1
        if self.pasos == 1:
            self.destino.write(self.contenido[:2])
            return None, False
        self.destino.write(self.contenido[2:])
        return None, True


def crear_pdf(*, con_texto=True) -> bytes:
    escritor = PdfWriter()
    pagina = escritor.add_blank_page(width=200, height=200)
    if con_texto:
        fuente = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        pagina[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): fuente}),
        })
        contenido = DecodedStreamObject()
        contenido.set_data(b"BT /F1 12 Tf 20 100 Td (Texto institucional) Tj ET")
        pagina[NameObject("/Contents")] = escritor._add_object(contenido)
    salida = BytesIO()
    escritor.write(salida)
    return salida.getvalue()


def crear_docx(*, con_texto=True) -> bytes:
    documento = Document()
    if con_texto:
        documento.add_paragraph("Introduccion institucional")
        documento.add_heading("Inscripciones", level=1)
        documento.add_paragraph("Consultar el calendario oficial.")
    salida = BytesIO()
    documento.save(salida)
    return salida.getvalue()


def test_listado_reune_paginas_y_escapa_la_carpeta(monkeypatch):
    servicio = ServicioFalso(paginas=[
        {"files": [{"id": "uno"}], "nextPageToken": "pagina-dos"},
        {"files": [{"id": "dos"}]},
    ])
    monkeypatch.setenv("GOOGLE_DRIVE_FOLDER_ID", "carpeta\\con'comilla")

    resultado = cliente.listar_archivos_institucionales(servicio=servicio, estricto=True)

    assert resultado == [{"id": "uno"}, {"id": "dos"}]
    primera, segunda = (llamada[1] for llamada in servicio.llamadas)
    assert primera["q"] == "'carpeta\\\\con\\'comilla' in parents and trashed=false"
    assert "nextPageToken" in primera["fields"]
    assert "pageToken" not in primera
    assert segunda["pageToken"] == "pagina-dos"
    assert servicio.cierres == 0  # El cliente inyectado pertenece a quien lo llamó.


def test_una_carpeta_vacia_es_resultado_valido_en_modo_estricto():
    assert cliente.listar_archivos_institucionales(
        servicio=ServicioFalso(), estricto=True
    ) == []


def test_error_de_paginacion_no_devuelve_resultados_parciales():
    servicio = ServicioFalso(paginas=[
        {"files": [{"id": "uno"}], "nextPageToken": "otra"},
        RuntimeError("contenido-privado"),
    ])
    with pytest.raises(ErrorDrive, match="listar") as error:
        cliente.listar_archivos_institucionales(servicio=servicio, estricto=True)
    assert "contenido-privado" not in str(error.value)
    assert error.value.__cause__ is None


def test_token_de_pagina_repetido_se_rechaza():
    servicio = ServicioFalso(paginas=[
        {"files": [], "nextPageToken": "repetida"},
        {"files": [], "nextPageToken": "repetida"},
    ])
    with pytest.raises(ErrorDrive):
        cliente.listar_archivos_institucionales(servicio=servicio, estricto=True)


@pytest.mark.parametrize("falla", [False, True])
def test_se_cierra_el_cliente_creado_localmente(monkeypatch, falla):
    servicio = ServicioFalso(paginas=[RuntimeError("no mostrar")] if falla else None)
    monkeypatch.setattr(cliente, "conectar_drive", lambda: servicio)

    cliente.listar_archivos_institucionales()

    assert servicio.cierres == 1


def test_error_legado_y_conexion_no_exponen_payload(monkeypatch, capsys):
    servicio = ServicioFalso(paginas=[RuntimeError("CLAVE_FICTICIA texto privado")])
    assert cliente.listar_archivos_institucionales(servicio=servicio) == []
    monkeypatch.setattr(cliente, "conectar_drive", lambda: None)
    with pytest.raises(ErrorDrive):
        cliente.listar_archivos_institucionales(estricto=True)
    assert "CLAVE_FICTICIA" not in capsys.readouterr().out


def test_metadatos_por_id_piden_solo_los_campos_necesarios():
    servicio = ServicioFalso()

    resultado = cliente.obtener_metadatos_archivo("documento-ficticio", servicio=servicio)

    assert resultado == servicio.metadatos
    assert servicio.llamadas == [("get", {
        "fileId": "documento-ficticio", "fields": "id,name,mimeType,webViewLink"
    })]
    assert servicio.cierres == 0


def test_metadatos_fallidos_no_exponen_la_respuesta():
    servicio = ServicioFalso(metadatos=RuntimeError("CLAVE_FICTICIA texto privado"))
    with pytest.raises(ErrorDrive) as error:
        cliente.obtener_metadatos_archivo("documento-ficticio", servicio=servicio)
    assert "CLAVE_FICTICIA" not in str(error.value)
    assert error.value.__cause__ is None


@pytest.mark.parametrize("tipo_mime, metodo", [
    (None, "get_media"),
    ("application/pdf", "get_media"),
    (cliente.TIPO_GOOGLE_DOCS, "export_media"),
])
def test_descarga_y_exportacion_son_bytes_en_memoria(monkeypatch, tipo_mime, metodo):
    monkeypatch.setattr(cliente, "MediaIoBaseDownload", DescargaFalsa)
    servicio = ServicioFalso(contenido=b"contenido institucional completo")

    resultado = cliente.descargar_contenido_archivo(
        "documento-ficticio", tipo_mime=tipo_mime, servicio=servicio, estricto=True
    )

    assert resultado == b"contenido institucional completo"
    esperado = {"fileId": "documento-ficticio"}
    if metodo == "export_media":
        esperado["mimeType"] = "application/pdf"
    assert servicio.llamadas == [(metodo, esperado)]
    assert servicio.cierres == 0


def test_descarga_fallida_es_segura_y_cierra_el_cliente(monkeypatch, capsys):
    monkeypatch.setattr(cliente, "MediaIoBaseDownload", DescargaFalsa)
    servicio = ServicioFalso(contenido=RuntimeError("CLAVE_FICTICIA texto privado"))
    monkeypatch.setattr(cliente, "conectar_drive", lambda: servicio)

    with pytest.raises(ErrorDrive) as error:
        cliente.descargar_contenido_archivo("documento-ficticio", estricto=True)
    assert servicio.cierres == 1
    assert "CLAVE_FICTICIA" not in str(error.value)
    assert error.value.__cause__ is None
    assert cliente.descargar_contenido_archivo("documento-ficticio") is None
    assert "CLAVE_FICTICIA" not in capsys.readouterr().out


def test_extrae_pdf_real_desde_bytes():
    secciones = extraer_texto_desde_bytes(crear_pdf(), "Institucional.pdf")
    assert secciones == [{"section": "General", "content": "Texto institucional"}]


def test_google_docs_exportado_usa_el_extractor_pdf():
    secciones = extraer_texto_desde_bytes(
        crear_pdf(), "Preguntas_frecuentes", tipo_mime=cliente.TIPO_GOOGLE_DOCS
    )
    assert secciones[0]["content"] == "Texto institucional"


def test_extrae_docx_real_por_secciones_desde_bytes():
    assert extraer_texto_desde_bytes(crear_docx(), "Documento.DOCX") == [
        {"section": "General", "content": "Introduccion institucional"},
        {"section": "Inscripciones", "content": "Consultar el calendario oficial."},
    ]


@pytest.mark.parametrize("contenido,nombre", [
    (b"", "vacio.pdf"),
    (b"contenido privado", "desconocido.txt"),
    (b"CLAVE_FICTICIA texto privado", "ilegible.pdf"),
    (b"CLAVE_FICTICIA texto privado", "ilegible.docx"),
])
def test_errores_de_extraccion_no_se_convierten_en_texto(contenido, nombre, caplog, capsys):
    with pytest.raises(ErrorExtraccion) as error:
        extraer_texto_desde_bytes(contenido, nombre)
    assert "CLAVE_FICTICIA" not in str(error.value)
    assert "CLAVE_FICTICIA" not in caplog.text
    salida = capsys.readouterr()
    assert "CLAVE_FICTICIA" not in salida.out + salida.err


@pytest.mark.parametrize("contenido,nombre", [
    (lambda: crear_pdf(con_texto=False), "sin_texto.pdf"),
    (lambda: crear_docx(con_texto=False), "sin_texto.docx"),
])
def test_documentos_sin_texto_se_rechazan(contenido, nombre):
    with pytest.raises(ErrorExtraccion, match="texto extraíble"):
        extraer_texto_desde_bytes(contenido(), nombre)


def test_interfaz_legada_reutiliza_extraccion_y_conserva_error(tmp_path):
    archivo = tmp_path / "institucional.docx"
    archivo.write_bytes(crear_docx())
    assert extractores.extract_text(str(archivo)) == extraer_texto_desde_bytes(
        archivo.read_bytes(), archivo.name
    )
    error = extractores.extract_text(str(tmp_path / "inexistente.pdf"))
    assert error[0]["section"] == "Error"
